from openpyxl import load_workbook


def parse_xlsx(path: str) -> str:
    wb = load_workbook(path, data_only=True)
    rows = []
    for ws in wb.worksheets:
        rows.append(f"# Sheet: {ws.title}")
        for row in ws.iter_rows(values_only=True):
            rows.append(" | ".join("" if v is None else str(v) for v in row))
    return "\n".join(rows)
